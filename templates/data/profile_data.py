"""표 데이터 첫 점검 (CSV·TSV·XLSX). 표준 라이브러리만 쓴다 (설치 없음).

실행: python3 profile_data.py <파일> [<파일> ...] [--by 열[,열]] [--max-rows 500000]
출력: 행·열 · 완전 중복 행 · 열마다 형식·빈 칸·값 종류 · 숫자 범위 · 확인할 것(질문 후보)
확인할 것: 표기만 다른 값(Etch01·ETCH-01·degC=C) · 줄임 표기(L1·line-1) · 시간대·날짜 형식 섞임(일/월/연 포함)
           · 단위 섞임 · 빈 칸 · 숫자 열 속 글자 · 극단값
--by 열[,열]: 그 열 값마다 숫자 열의 최소·평균·최대 (예: --by equipment,sensor 로 설비·센서별 단위 차이 찾기)
           --by 가 없고 200행 이상이면 값 종류 2~12개인 글자 열(두 열이면 종류 곱 24 이하)로 자동으로 묶는다
큰 파일도 한 줄씩 읽는다. 값 종류는 열마다 2만 개까지만 센다.
"""
import csv
import math
import re
import sys
import zipfile
from collections import Counter, defaultdict
from xml.etree import ElementTree as ET

NULLS = {"", "na", "n/a", "nan", "null", "none", "-", "--", "?", "없음", "#n/a"}
NUM_RE = re.compile(r"^[+-]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?([eE][+-]?\d+)?$|^[+-]?\.\d+$")
UNIT_RE = re.compile(r"^[+-]?\d[\d,]*(\.\d+)?\s*([A-Za-z%°µ/가-힣]+[A-Za-z0-9%°/]*)$")
DATE_RE = re.compile(r"^(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{4})([ T]\d{1,2}:\d{2}(:\d{2})?(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?$")
TZ_RE = re.compile(r"(Z|[+-]\d{2}:?\d{2})$")
DMY_RE = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.]\d{4}")
KEEP = 20000
TAB = chr(9)


def norm(v):
    v = re.sub(r"(?i)^deg(?=[cf]$)", "", v)
    return re.sub(r"[\s\-_.·°/]+", "", v).casefold()


def abbr_of(short, long):  # short 의 글자가 long 안에 차례대로 있나 (l · line)
    it = iter(long)
    return all(ch in it for ch in short)


def to_num(v):
    return float(v.replace(",", "")) if NUM_RE.match(v) else None


def read_csv(path):
    head = open(path, "rb").read(1 << 20)
    enc = "latin-1"
    for e in ("utf-8-sig", "cp949"):
        try:
            head.decode(e)
            enc = e
            break
        except UnicodeDecodeError as err:
            if err.start >= len(head) - 3:  # 1MB 경계에서 잘린 글자
                enc = e
                break
    sample = head[:20000].decode(enc, errors="replace")
    sample = sample[: sample.rfind("\n") + 1] or sample
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="," + TAB + ";|")
    except csv.Error:
        dialect = csv.excel_tab if path.lower().endswith(".tsv") else csv.excel
    rows = csv.reader(open(path, encoding=enc, errors="replace", newline=""), dialect)
    return enc, next(rows, []), rows


def read_xlsx(path, max_rows):
    """첫 시트만. openpyxl 이 있으면 쓰고, 없으면 zip·xml 로 직접 읽는다 (날짜는 숫자로 보임)."""
    try:
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = ["" if c is None else str(c) for c in next(it, [])]
        rows = (["" if c is None else str(c) for c in r] for r in it)
        return "openpyxl", header, rows
    except ImportError:
        pass
    m = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    z = zipfile.ZipFile(path)
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        for si in ET.fromstring(z.read("xl/sharedStrings.xml")).iter(m + "si"):
            shared.append("".join(t.text or "" for t in si.iter(m + "t")))
    sheet = sorted(n for n in z.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n))[0]
    table = []
    for row in ET.fromstring(z.read(sheet)).iter(m + "row"):
        cells = {}
        for c in row.iter(m + "c"):
            col = 0
            for ch in re.match(r"[A-Z]*", c.get("r", "A")).group() or "A":
                col = col * 26 + ord(ch) - 64
            t, v = c.get("t"), c.find(m + "v")
            if t == "s" and v is not None:
                val = shared[int(v.text)]
            elif t == "inlineStr":
                val = "".join(x.text or "" for x in c.iter(m + "t"))
            else:
                val = v.text if v is not None and v.text else ""
            cells[col - 1] = val
        if cells:
            table.append([cells.get(i, "") for i in range(max(cells) + 1)])
        if len(table) > max_rows:
            break
    return "xlsx(zip)", (table[0] if table else []), iter(table[1:])


class Col:
    def __init__(self):
        self.n = self.empty = self.nums = self.dates = 0
        self.null_words, self.vals = Counter(), Counter()
        self.overflow = False
        self.mean = self.m2 = 0.0
        self.lo, self.hi = math.inf, -math.inf
        self.units, self.shapes, self.tz, self.text_in_num = Counter(), Counter(), Counter(), Counter()
        self.dmy = [0, 0, 0]  # 연도가 끝에 오는 날짜: 앞자리 최대 · 둘째 최대 · 개수

    def add(self, v):
        self.n += 1
        s = v.strip()
        if s.casefold() in NULLS:
            self.empty += 1
            if s:
                self.null_words[s] += 1
            return
        if s in self.vals or len(self.vals) < KEEP:
            self.vals[s] += 1
        else:
            self.overflow = True
        x = to_num(s)
        if x is not None:
            self.nums += 1
            d = x - self.mean
            self.mean += d / self.nums
            self.m2 += d * (x - self.mean)
            self.lo, self.hi = min(self.lo, x), max(self.hi, x)
            return
        if DATE_RE.match(s):
            self.dates += 1
            self.shapes[re.sub(r"\d", "9", s)] += 1
            z = TZ_RE.search(s[10:])
            self.tz[z.group(1) if z else "없음"] += 1
            m = DMY_RE.match(s)
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                self.dmy = [max(self.dmy[0], a), max(self.dmy[1], b), self.dmy[2] + 1]
            return
        u = UNIT_RE.match(s)
        if u:
            self.units[u.group(2)] += 1
        if len(self.text_in_num) < 50:
            self.text_in_num[s] += 1

    def kind(self):
        filled = self.n - self.empty
        if filled and self.nums / filled >= 0.9:
            return "숫자"
        if filled and self.dates / filled >= 0.9:
            return "날짜"
        return "글자"


def fmt(x):
    if math.isinf(x):
        return "-"
    return f"{x:.4g}" if abs(x) < 1e6 else f"{x:.3e}"


def open_table(path, max_rows):
    if path.lower().endswith((".xlsx", ".xlsm")):
        enc, header, rows = read_xlsx(path, max_rows)
    else:
        enc, header, rows = read_csv(path)
    return enc, [h.strip() or f"열{i + 1}" for i, h in enumerate(header)], rows


def add_group(stat, r, keys, header):
    g = " · ".join(r[k].strip() if k < len(r) else "" for k in keys)
    for i, v in enumerate(r[: len(header)]):
        x = None if i in keys else to_num(v.strip())
        if x is not None:
            st = stat[(g, header[i])]
            st[0], st[1], st[2], st[3] = min(st[0], x), max(st[1], x), st[2] + x, st[3] + 1


def profile(path, by=None, max_rows=500000):
    enc, header, rows = open_table(path, max_rows)
    cols = [Col() for _ in header]
    keys = [header.index(k) for k in (by or "").split(",") if k in header]
    stat = defaultdict(lambda: [math.inf, -math.inf, 0.0, 0])
    seen, dup, width_bad, n = set(), 0, 0, 0
    for r in rows:
        if n >= max_rows:
            break
        n += 1
        if len(r) != len(header):
            width_bad += 1
        h = hash(tuple(r))
        if h in seen:
            dup += 1
        else:
            seen.add(h)
        for i, c in enumerate(cols):
            c.add(r[i] if i < len(r) else "")
        if keys:
            add_group(stat, r, keys, header)
    print(f"\n## {path}")
    print(f"- 읽기 {enc} · 행 {n}{' (앞부분만, --max-rows)' if n >= max_rows else ''} · 열 {len(header)}"
          f" · 완전 중복 행 {dup}" + (f" · 열 수가 다른 행 {width_bad}" if width_bad else ""))
    checks = [f"완전 중복 행 {dup}개 · 재전송·중복 기록인지, 하나로 볼지"] if dup else []
    print("\n| 열 | 형식 | 빈 칸 | 값 종류 | 범위·자주 나오는 값 |\n|---|---|---|---|---|")
    for name, c in zip(header, cols):
        kinds, k = len(c.vals), c.kind()
        if k == "숫자":
            sd = math.sqrt(c.m2 / c.nums) if c.nums > 1 else 0.0
            desc = f"최소 {fmt(c.lo)} · 평균 {fmt(c.mean)} · 최대 {fmt(c.hi)}"
            if sd and (c.hi - c.mean > 6 * sd or c.mean - c.lo > 6 * sd):
                checks.append(f"{name}: 평균±6표준편차 밖 값 있음 (최소 {fmt(c.lo)} · 최대 {fmt(c.hi)}) · 단위·오기·실제 이상인지")
            if c.text_in_num:
                checks.append(f"{name}: 숫자 열에 글자 값 {list(c.text_in_num)[:4]}")
        else:
            top = kinds if k == "글자" and kinds <= 20 else (8 if kinds <= 8 else 5)
            desc = " · ".join(f"{v[:24]}({cnt})" for v, cnt in c.vals.most_common(top))
        empty = f"{c.empty}" + (f" (표기 {dict(c.null_words.most_common(3))})" if c.null_words else "")
        print(f"| {name} | {k} | {empty} | {kinds}{'+' if c.overflow else ''} | {desc} |")
        if k != "숫자" and 1 < kinds <= 2000:
            groups = defaultdict(list)
            for v in c.vals:
                groups[norm(v)].append(v)
            same = [g for g in groups.values() if len(g) > 1]
            if same:
                shown = " / ".join("=".join(g[:4]) for g in same[:5])
                checks.append(f"{name}: 표기만 다른 값 {len(same)}묶음 ({shown}) · 같은 대상인지")
            near = defaultdict(dict)  # 숫자가 같고 글자가 줄임말 관계인 값 (L1 · line-1)
            for v in c.vals:
                d = re.sub(r"\D", "", v).lstrip("0")
                if d and len(v) <= 12:
                    near[(v[:1].casefold(), d)].setdefault(norm(v), v)
            pairs = []
            for g in near.values():
                ks = sorted(g, key=len)
                al = [re.sub(r"[^a-z]", "", k) for k in ks]
                hit = [g[k] for k, a in zip(ks, al) if abbr_of(al[0], a)]
                if len(hit) > 1:
                    pairs.append(hit)
            if pairs:
                checks.append(f"{name}: 줄임 표기 {len(pairs)}묶음 ({' / '.join('='.join(h[:4]) for h in pairs[:5])}) · 같은 대상인지")
            if kinds <= 12 and re.search(r"unit|단위", name, re.I):
                checks.append(f"{name}: 단위 {kinds}가지 {list(c.vals)} · 환산이 필요한지, 언제부터 바뀌었는지")
        if c.empty and n >= 50:  # 작은 표는 에이전트가 직접 읽는다
            checks.append(f"{name}: 빈 칸 {c.empty}개 · 빈 칸의 뜻(미완료·끊김·0)과 계산에 넣을지")
        if len(c.shapes) > 1:
            checks.append(f"{name}: 날짜 형식 {len(c.shapes)}가지 {dict(c.shapes.most_common(3))}")
        if c.dmy[2]:
            a, b = c.dmy[0], c.dmy[1]
            order = "일/월/연(앞자리가 12 초과)" if a > 12 else "월/일/연(둘째가 12 초과)" if b > 12 else "일/월인지 월/일인지 모름(둘 다 12 이하)"
            checks.append(f"{name}: 연도가 끝에 오는 날짜 {c.dmy[2]}개 · {order} · 앞뒤 행으로 정할 것")
        if len(c.tz) > 1:
            checks.append(f"{name}: 시간대 표기 섞임 {dict(c.tz)} · 표기 없는 값은 어느 시간대인지")
        if len(c.units) > 1:
            checks.append(f"{name}: 값 속 단위 섞임 {dict(c.units.most_common(6))} · 바뀐 시점·환산 기준")
    if by and not keys:
        print(f"\n- --by {by}: 열을 찾지 못함 (열: {', '.join(header)})")
    limit, label = 80, by
    if not by and n >= 200 and any(c.kind() == "숫자" for c in cols):
        cand = [i for i, c in enumerate(cols) if c.kind() == "글자" and 2 <= len(c.vals) <= 12 and not c.overflow]
        kv = lambda ks: math.prod(len(cols[i].vals) for i in ks)  # noqa: E731
        opts = [[i, j] for i in cand for j in cand if i < j and kv([i, j]) <= 24] or [[i] for i in cand]
        keys = max(opts, key=kv) if opts else []
        if keys:
            limit, label = 24, ",".join(header[i] for i in keys) + " (자동)"
            _, _, rows2 = open_table(path, max_rows)
            for j, r in enumerate(rows2):
                if j >= max_rows:
                    break
                add_group(stat, r, keys, header)
    if keys and stat:
        print(f"\n| {label} | 숫자 열 | 최소 | 평균 | 최대 | 개수 |\n|---|---|---|---|---|---|")
        for (g, name), st in sorted(stat.items())[:limit]:
            print(f"| {g} | {name} | {fmt(st[0])} | {fmt(st[2] / st[3])} | {fmt(st[1])} | {st[3]} |")
        if len(stat) > limit:
            print(f"- 묶음 {len(stat)}줄 중 앞 {limit}줄 · 좁히려면 --by 열")
    print("\n### 확인할 것 (실무자 질문 후보)")
    for ch in checks or ["눈에 띄는 것 없음 · 빈 칸의 뜻과 기간 경계는 따로 확인"]:
        print(f"- {ch}")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    files, by, mx = [], None, 500000
    it = iter(sys.argv[1:])
    for a in it:
        if a == "--by":
            by = next(it, None)
        elif a == "--max-rows":
            mx = int(next(it, mx))
        else:
            files.append(a)
    if not files:
        print(__doc__)
        sys.exit(1)
    print("# 데이터 점검 (profile_data.py)")
    for p in files:
        try:
            profile(p, by, mx)
        except Exception as e:  # 한 파일이 깨져도 나머지는 본다
            print(f"\n## {p}\n- 읽기 실패: {type(e).__name__}: {e}")
